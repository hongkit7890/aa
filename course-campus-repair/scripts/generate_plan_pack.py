#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate remaining plan deliverables for Form 3 campus repair course."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
STARTER = ROOT / "starter-kit"
OUT.mkdir(parents=True, exist_ok=True)
STARTER.mkdir(parents=True, exist_ok=True)


def set_run_font(run, name="Microsoft JhengHei", size=11, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def style_doc(doc):
    s = doc.sections[0]
    s.top_margin = Cm(1.8)
    s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2)
    s.right_margin = Cm(2)


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=15 if level == 1 else 12, bold=True)
    return p


def p(doc, text, bold=False, size=11, after=4):
    para = doc.add_paragraph()
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold)
    para.paragraph_format.space_after = Pt(after)
    return para


def bullets(doc, items, size=10):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        run = para.add_run(item)
        set_run_font(run, size=size)
        para.paragraph_format.space_after = Pt(1)


def cell(c, text, bold=False, size=9):
    c.text = ""
    run = c.paragraphs[0].add_run(text)
    set_run_font(run, size=size, bold=bold)


def fill(table, rows):
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell(table.rows[i].cells[j], str(val), bold=(i == 0), size=9)


LESSONS = [
    {
        "no": 1,
        "title": "系統介紹與數據庫基礎",
        "concept": "資料表／欄位／記錄／主鍵",
        "outcome": "campus_system + repairs + 手動 3 筆資料",
        "mins": [
            ("導入：試用完成品", 15),
            ("概念：四格流程＋資料表像 Excel", 30),
            ("實作：建庫建表＋3筆＋截圖", 60),
            ("小結預告", 15),
        ],
        "teacher_tips": [
            "欄位用中文說明＋英文欄名對照",
            "強調 id 主鍵自動遞增",
            "少講完整架構，只要知道資料在資料庫",
        ],
        "notes": [
            "資料庫像練習簿；資料表像其中一頁表。",
            "欄位＝欄標題；記錄＝一橫列；主鍵＝永不重複編號。",
            "本課：資料庫 campus_system／資料表 repairs。",
            "四格流程：填表 → PHP → 資料庫 → 列表（今天先做資料庫）。",
        ],
        "steps": [
            "啟動 XAMPP（Apache + MySQL）",
            "phpMyAdmin 建立 campus_system",
            "建立 repairs，設定 id 為主鍵 A_I",
            "手動新增 3 筆報修",
            "截圖上傳",
        ],
        "checklist": [
            "能說出資料表／欄位／記錄／主鍵",
            "campus_system 與 repairs 已建立",
            "至少 3 筆資料並已截圖",
        ],
        "star": ["完成建表＋3筆（必做）", "多加 2 筆不同地點資料", "畫出四格流程並標註今天完成哪一格"],
        "peer": [
            "對方資料庫名稱是否正確？",
            "id 是否為主鍵且自動遞增？",
            "是否有 name／location／description／status？",
            "是否至少 3 筆？",
            "截圖是否清楚可見表格？",
        ],
    },
    {
        "no": 2,
        "title": "PHP基礎與表單處理",
        "concept": "表單 + $_POST",
        "outcome": "form.html + show.php 可運行（尚未寫入資料庫）",
        "mins": [
            ("回顧資料表欄位", 10),
            ("概念：PHP 極簡＋表單＋$_POST", 30),
            ("實作：form.html + show.php", 65),
            ("小結預告", 15),
        ],
        "teacher_tips": [
            "本堂不教 for、不連資料庫",
            "show.php 用大字顯示成功感",
            "錯誤三招：副檔名、路徑、大小寫",
        ],
        "notes": [
            "PHP 寫在 <?php ... ?>。",
            "變數以 $ 開頭；用 echo 顯示。",
            "表單：method=\"post\"、action=\"show.php\"。",
            "接收：$name = $_POST['name'];",
        ],
        "steps": [
            "在 htdocs 建立 campus_repair 資料夾",
            "建立 form.html（姓名、地點、描述）",
            "建立 show.php 大字顯示三欄",
            "瀏覽器測試送出至少 2 次",
            "卡住用 var_dump($_POST)",
        ],
        "checklist": [
            "表單可以打開",
            "送出後顯示正確內容",
            "知道 $_POST 的用途",
        ],
        "star": ["表單＋顯示頁可運行", "加入簡單 if 空值提示", "美化顯示字級／顏色（不加複雜CSS也可）"],
        "peer": [
            "input 的 name 是否與 $_POST 一致？",
            "action 是否指向 show.php？",
            "method 是否為 post？",
            "三個欄位都能顯示？",
            "檔名大小寫是否正確？",
        ],
    },
    {
        "no": 3,
        "title": "PHP連接數據庫與新增資料",
        "concept": "連線 + INSERT",
        "outcome": "表單資料成功寫入 repairs",
        "mins": [
            ("回顧欄位對應", 10),
            ("概念：db.php、INSERT、安全一句", 30),
            ("實作：insert.php 寫入＋核對", 65),
            ("小結預告", 15),
        ],
        "teacher_tips": [
            "發 db.php，學生 require 即可",
            "prepared statement 示範一次即可",
            "預留故意寫錯再修好",
        ],
        "notes": [
            "require 'db.php'; 使用老師連線。",
            "INSERT 把表單資料寫進 repairs。",
            "安全一句：不要把輸入直接拼進 SQL。",
            "成功後到 phpMyAdmin 刷新看新列。",
        ],
        "steps": [
            "複製 Starter Kit 的 db.php",
            "表單 action 改 insert.php",
            "完成 insert.php 填空（或使用範本）",
            "送出後在 phpMyAdmin 確認",
            "加分：成功後跳轉 list.php",
        ],
        "checklist": [
            "表單指向 insert.php",
            "資料庫多一筆新資料",
            "能指出 INSERT 在做什麼",
        ],
        "star": ["寫入資料庫", "寫入後跳轉 list.php", "空欄位不允許送出（簡單 if）"],
        "peer": [
            "db.php 是否在正確資料夾？",
            "MySQL 是否已啟動？",
            "INSERT 欄位是否對應表單？",
            "phpMyAdmin 是否看到新列？",
            "status 是否為 pending？",
        ],
    },
    {
        "no": 4,
        "title": "查詢與顯示資料",
        "concept": "SELECT + 迴圈出表格",
        "outcome": "list.php 管理頁＋狀態篩選",
        "mins": [
            ("回顧 INSERT", 10),
            ("概念：SELECT、while 表格、篩選", 30),
            ("實作：list.php＋篩選", 65),
            ("小結預告", 15),
        ],
        "teacher_tips": [
            "列表頁時間可略多——最有系統感",
            "用 PHP開→表格→PHP關 三段式",
            "篩選保持下拉簡單",
        ],
        "notes": [
            "SELECT * FROM repairs; 取出全部。",
            "while 迴圈一列輸出一個 <tr>。",
            "篩選：WHERE status='pending'／'done'。",
            "口訣：PHP 開 → 表格 → PHP 關。",
        ],
        "steps": [
            "建立／完成 list.php",
            "查詢並以表格顯示全部報修",
            "加入狀態下拉篩選",
            "測試三種篩選",
            "加分：依時間新到舊排序",
        ],
        "checklist": [
            "列表能顯示資料庫資料",
            "篩選可以使用",
            "表格欄位齊全易讀",
        ],
        "star": ["全部列表", "狀態篩選", "ORDER BY created_at DESC"],
        "peer": [
            "表格是否有資料？",
            "欄位是否齊全？",
            "篩選待處理是否有效？",
            "篩選已完成是否有效？",
            "中文是否正常顯示？",
        ],
    },
    {
        "no": 5,
        "title": "功能整合、更新狀態與成果發表",
        "concept": "UPDATE + 整合展示",
        "outcome": "可展示完整報修小系統＋反思",
        "mins": [
            ("回顧新增與列表", 10),
            ("概念：UPDATE、導覽、現成CSS", 25),
            ("實作：update＋整合＋美化", 55),
            ("成果發表（每組2分鐘）", 20),
            ("倫理反思三句", 10),
        ],
        "teacher_tips": [
            "CSS 給現成檔，只改標題／主色",
            "UPDATE 用 update.php?id=&status=",
            "展示流程：送出→列表→改狀態",
        ],
        "notes": [
            "UPDATE repairs SET status='done' WHERE id=?;",
            "列表加連結：update.php?id=1&status=done",
            "套用 style.css，可讀比花俏重要。",
            "倫理：少敏感資料；不上公開網；尊重內容。",
        ],
        "steps": [
            "完成 update.php",
            "list.php 每列加「標示完成」",
            "整合導覽：表單⇄列表",
            "套用 style.css",
            "展示2分鐘＋反思三題",
        ],
        "checklist": [
            "可以更新狀態",
            "頁面能互相跳轉",
            "完成展示與反思",
        ],
        "star": ["UPDATE 狀態", "完整導覽整合", "延伸構想（借用／問卷）一頁說明"],
        "peer": [
            "能否把 pending 改 done？",
            "改完列表是否更新？",
            "表單與列表能否互跳？",
            "展示是否含送出→列表→改狀態？",
            "反思三題是否填寫？",
        ],
    },
]


# ---------------------------------------------------------------------------
# 1) Refined teacher lesson plans (compact)
# ---------------------------------------------------------------------------

def build_teacher_plans():
    doc = Document()
    style_doc(doc)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("精簡教師教案（五堂）")
    set_run_font(r, size=16, bold=True)
    p(doc, "課程：校園小系統實作：數據庫與PHP入門｜對象：中三｜專題：校園報修（全班統一）", size=10)
    p(doc, "原則：每堂只學 1 個新概念 + 1 個可見成果；先求有再求好；提供 Starter Kit。", size=10, bold=True)

    for lesson in LESSONS:
        if lesson["no"] > 1:
            doc.add_page_break()
        h(doc, f"第{lesson['no']}堂：{lesson['title']}（120分鐘）")
        p(doc, f"唯一新概念：{lesson['concept']}", bold=True)
        p(doc, f"下課前必見成果：{lesson['outcome']}", bold=True)

        table = doc.add_table(rows=len(lesson["mins"]) + 1, cols=3)
        table.style = "Table Grid"
        cell(table.rows[0].cells[0], "階段", True)
        cell(table.rows[0].cells[1], "分鐘", True)
        cell(table.rows[0].cells[2], "教師要注意", True)
        tips = lesson["teacher_tips"] + [""] * 5
        for i, (stage, m) in enumerate(lesson["mins"], start=1):
            cell(table.rows[i].cells[0], stage)
            cell(table.rows[i].cells[1], str(m))
            cell(table.rows[i].cells[2], tips[i - 1] if i - 1 < len(lesson["teacher_tips"]) else "巡堂協助")

        p(doc, "分層任務", bold=True)
        bullets(doc, [f"★ {lesson['star'][0]}", f"★★ {lesson['star'][1]}", f"★★★ {lesson['star'][2]}"])
        p(doc, "同儕互查（5問）", bold=True)
        bullets(doc, [f"{i}. {q}" for i, q in enumerate(lesson["peer"], 1)])
        p(doc, "成功畫面：學生應看到 —— " + lesson["outcome"], size=10)

    path = OUT / "05_精簡教師教案_五堂.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 2) Half-page notes + task sheets × 5
# ---------------------------------------------------------------------------

def build_notes_and_tasks():
    # Combined printable pack
    doc = Document()
    style_doc(doc)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("學生半頁筆記＋任務單＋自我檢查（五堂）")
    set_run_font(r, size=16, bold=True)
    p(doc, "姓名：________　班別：____　學號：________　組別：____", size=10)

    for lesson in LESSONS:
        doc.add_page_break()
        h(doc, f"第{lesson['no']}堂筆記｜{lesson['title']}")
        p(doc, f"今天只學：{lesson['concept']}", bold=True, size=11)
        p(doc, f"成功長這樣：{lesson['outcome']}", bold=True, size=11)
        p(doc, "【半頁筆記】", bold=True)
        bullets(doc, lesson["notes"], size=10)

        p(doc, "【任務步驟】", bold=True)
        for i, step in enumerate(lesson["steps"], 1):
            p(doc, f"{i}. {step}", size=10)

        p(doc, "【自我檢查】", bold=True)
        for c in lesson["checklist"]:
            p(doc, f"☐ {c}", size=10)

        p(doc, "【分層】", bold=True)
        p(doc, f"★ {lesson['star'][0]}　★★ {lesson['star'][1]}　★★★ {lesson['star'][2]}", size=9)

        p(doc, "【同儕互查】對方姓名：________", bold=True, size=10)
        for i, q in enumerate(lesson["peer"], 1):
            p(doc, f"{i}. {q}　☐是 ☐否", size=9)

        p(doc, "教師簽署：________　日期：________", size=9)

    path = OUT / "06_學生半頁筆記與任務單_五堂.docx"
    doc.save(path)

    # Also separate one-pagers for convenience
    paths = [path]
    for lesson in LESSONS:
        d = Document()
        style_doc(d)
        h(d, f"第{lesson['no']}堂｜半頁筆記＋任務單")
        p(d, f"概念：{lesson['concept']}｜成果：{lesson['outcome']}", bold=True, size=10)
        p(d, "筆記", bold=True)
        bullets(d, lesson["notes"])
        p(d, "步驟", bold=True)
        for i, step in enumerate(lesson["steps"], 1):
            p(d, f"{i}. {step}", size=10)
        p(d, "自我檢查", bold=True)
        for citem in lesson["checklist"]:
            p(d, f"☐ {citem}", size=10)
        sp = OUT / f"06_第{lesson['no']}堂_筆記任務單.docx"
        d.save(sp)
        paths.append(sp)
    return paths


# ---------------------------------------------------------------------------
# 3) Rubric + cheat sheet + error table
# ---------------------------------------------------------------------------

def build_rubric_and_refs():
    doc = Document()
    style_doc(doc)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("評分規準＋作弊條＋錯誤對照表")
    set_run_font(r, size=16, bold=True)

    h(doc, "一、評分規準表（Rubric）")
    p(doc, "總結性以「能不能跑通」為主；加分項另計，避免落後者焦慮。", size=10)
    rub = doc.add_table(rows=7, cols=5)
    rub.style = "Table Grid"
    fill(
        rub,
        [
            ("評分項目", "優異", "良好", "合格", "待改進"),
            (
                "資料庫設計",
                "欄位完整、主鍵正確、能清楚說明",
                "欄位正確、說明大致清楚",
                "能建表並有資料",
                "未完成或錯誤多",
            ),
            (
                "PHP程式邏輯",
                "流程順暢、能解釋關鍵碼",
                "可運行、解釋大致正確",
                "在協助下可運行",
                "多數無法運行",
            ),
            (
                "系統功能完整度",
                "新增／查詢／更新齊備穩定",
                "三項皆有、偶有小問題",
                "完成新增＋查詢",
                "僅部分頁面",
            ),
            (
                "除錯與問題解決",
                "能獨立對照錯誤並修正",
                "多數能自行修正",
                "需提示才能修正",
                "遇錯即停",
            ),
            (
                "合作與展示",
                "展示清晰、分工明確、能答問",
                "展示清楚、能說明流程",
                "完成基本展示",
                "展示不清／未參與",
            ),
            (
                "資訊倫理與反思",
                "能連結隱私／安全／責任",
                "能說出主要倫理重點",
                "完成反思題",
                "空泛或未交",
            ),
        ],
    )

    p(doc, "可觀察通過標準（評分時勾選）", bold=True, size=10)
    bullets(
        doc,
        [
            "CP1：資料表存在且 ≥3 筆",
            "CP2：表單送出後網頁顯示正確欄位",
            "CP3：新資料出現在資料庫",
            "CP4：列表可顯示並可篩選",
            "CP5：可更新狀態並完成展示／反思",
        ],
        size=10,
    )
    p(doc, "加分（明確另計）：空欄位驗證、寫入後跳轉、統計／搜尋、版面美化、登入構想", size=10)

    h(doc, "二、一頁作弊條（Cheat Sheet）")
    p(doc, "PHP", bold=True)
    bullets(
        doc,
        [
            "<?php  $x = 1;  echo $x;  if ($x > 0) { echo 'ok'; }  ?>",
            "接收表單：$name = $_POST['name'];",
            "連線：require 'db.php';",
            "輸出：echo htmlspecialchars($name);",
        ],
        size=10,
    )
    p(doc, "SQL", bold=True)
    bullets(
        doc,
        [
            "INSERT INTO repairs (name, location, description, status) VALUES (?, ?, ?, 'pending');",
            "SELECT * FROM repairs;",
            "SELECT * FROM repairs WHERE status='pending';",
            "UPDATE repairs SET status='done' WHERE id=?;",
        ],
        size=10,
    )
    p(doc, "檔案地圖：form.html → insert.php → MySQL → list.php → update.php", bold=True, size=10)

    h(doc, "三、常見錯誤對照表")
    et = doc.add_table(rows=8, cols=2)
    et.style = "Table Grid"
    fill(
        et,
        [
            ("現象", "先檢查"),
            ("網頁打不開", "Apache 是否啟動；網址／資料夾路徑"),
            ("送出後空白或 Notice", "input name 是否與 $_POST 鍵名一致"),
            ("連線失敗", "MySQL 是否啟動；db.php 帳密與資料庫名"),
            ("INSERT 後看不到", "是否選錯資料庫；SQL／執行是否成功"),
            ("列表空白", "表內是否有資料；表名／欄名是否寫錯"),
            ("中文亂碼", "檔案與資料庫是否 UTF-8／utf8mb4"),
            ("更新沒反應", "id 參數是否傳到；UPDATE 後是否回到 list"),
        ],
    )

    h(doc, "四、延伸構想（加分／學有餘力）")
    bullets(
        doc,
        [
            "物品借用系統：物品名稱、借用人、應還日期、狀態",
            "校園問卷：題目、選項、提交時間；可做簡單統計",
            "提醒：仍建議先跑通報修主線再延伸",
        ],
        size=10,
    )

    path = OUT / "07_評分規準_作弊條_錯誤對照.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 4) PPT outline (markdown + docx)
# ---------------------------------------------------------------------------

PPT_OUTLINE = [
    ("封面", ["課程名稱", "中三｜5×2小時｜校園報修", "含 VASK"]),
    ("課程目標", ["數據庫概念", "基本 PHP", "完成小系統", "完整流程", "除錯與倫理"]),
    ("VASK 總覽", ["V 價值觀", "A 態度", "S 技能", "K 知識"]),
    ("系統四步", ["填表", "PHP", "資料庫", "列表"]),
    ("第1堂分隔頁", ["焦點：資料表概念", "成果：3筆資料"]),
    ("第1堂目標與VASK", ["建庫建表", "關懷校園", "主鍵"]),
    ("先玩完成品", ["送出一筆", "問資料去哪", "建立動機"]),
    ("四格流程", ["今天先攻資料庫"]),
    ("資料表像Excel", ["欄位／記錄／主鍵"]),
    ("欄位設計表", ["id name location description status created_at"]),
    ("第1堂實作", ["XAMPP", "建庫", "建表", "3筆", "截圖"]),
    ("第1堂檢查", ["檢查清單", "預告表單"]),
    ("第2堂分隔頁", ["焦點：$_POST", "成果：顯示頁"]),
    ("第2堂目標與VASK", ["極簡PHP", "先求有"]),
    ("PHP幾句", ["標籤、變數、echo、if"]),
    ("表單三要素", ["method action name"]),
    ("$_POST用法", ["三欄接收與echo"]),
    ("第2堂實作", ["form.html", "show.php"]),
    ("錯誤三招", ["副檔名、路徑、大小寫"]),
    ("第2堂檢查", ["尚未寫入資料庫"]),
    ("第3堂分隔頁", ["焦點：INSERT", "成果：進資料庫"]),
    ("第3堂目標與VASK", ["連線", "安全意識"]),
    ("資料流對照", ["表單欄↔資料表欄"]),
    ("db.php", ["require即可"]),
    ("INSERT概念", ["VALUES (?,?,?)"]),
    ("安全一句", ["勿直接拼接", "prepared statement"]),
    ("第3堂實作", ["insert.php", "phpMyAdmin核對"]),
    ("除錯示範", ["寫錯→讀錯→修好"]),
    ("第4堂分隔頁", ["焦點：SELECT", "成果：list.php"]),
    ("第4堂目標與VASK", ["迴圈出表格", "尊重內容"]),
    ("SELECT", ["全部／WHERE"]),
    ("迴圈出表格", ["PHP開表格PHP關"]),
    ("簡單篩選", ["下拉三種狀態"]),
    ("第4堂實作", ["列表＋篩選"]),
    ("第4堂檢查", ["預告UPDATE"]),
    ("第5堂分隔頁", ["焦點：UPDATE", "成果：可展示系統"]),
    ("第5堂目標與VASK", ["整合", "倫理"]),
    ("UPDATE狀態", ["連結帶id與status"]),
    ("整合清單", ["form insert list update"]),
    ("美化原則", ["現成CSS", "可讀優先"]),
    ("展示2分鐘", ["送出→列表→改狀態"]),
    ("倫理三句", ["敏感資料", "不上公開網", "尊重"]),
    ("課程總結", ["小型系統基本循環"]),
    ("附件與延伸", ["Starter Kit", "加分", "借用／問卷"]),
]


def build_ppt_outline():
    # Markdown
    md_lines = [
        "# 五堂教學簡報大綱／逐頁要點",
        "",
        "> 對應 `04_五堂教學簡報_合訂.pptx`（約 40–50 頁精簡版，非長篇理論）",
        "",
    ]
    for i, (title, points) in enumerate(PPT_OUTLINE, 1):
        md_lines.append(f"## 第 {i} 頁：{title}")
        for pt in points:
            md_lines.append(f"- {pt}")
        md_lines.append("")
    md_path = OUT / "08_PPT大綱與逐頁要點.md"
    md_path.write_text("\n".join(md_lines), encoding="utf-8")

    doc = Document()
    style_doc(doc)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("PPT 大綱與逐頁要點")
    set_run_font(r, size=16, bold=True)
    p(doc, f"共 {len(PPT_OUTLINE)} 頁要點｜對應合訂簡報與分堂簡報", size=10)
    for i, (title, points) in enumerate(PPT_OUTLINE, 1):
        p(doc, f"第{i}頁　{title}", bold=True, size=11)
        bullets(doc, points, size=10)

    docx_path = OUT / "08_PPT大綱與逐頁要點.docx"
    doc.save(docx_path)
    return [md_path, docx_path]


# ---------------------------------------------------------------------------
# 5) Starter Kit
# ---------------------------------------------------------------------------

def write_starter_kit():
    files = {}

    files["sql/schema.sql"] = """-- 校園報修系統：資料庫與資料表
-- 在 phpMyAdmin 可直接匯入，或手動執行

CREATE DATABASE IF NOT EXISTS campus_system
  DEFAULT CHARACTER SET utf8mb4
  COLLATE utf8mb4_general_ci;

USE campus_system;

CREATE TABLE IF NOT EXISTS repairs (
  id INT AUTO_INCREMENT PRIMARY KEY COMMENT '編號（主鍵）',
  name VARCHAR(50) NOT NULL COMMENT '報修人',
  location VARCHAR(100) NOT NULL COMMENT '地點',
  description TEXT NOT NULL COMMENT '問題描述',
  status VARCHAR(20) NOT NULL DEFAULT 'pending' COMMENT '狀態 pending/done',
  created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP COMMENT '建立時間'
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;

-- 範例資料（可選）
INSERT INTO repairs (name, location, description, status) VALUES
('陳同學', '3樓電腦室', '投影機無法開機', 'pending'),
('李老師', '圖書館', '冷氣漏水', 'pending'),
('王同學', '操場旁洗手間', '水龍頭損壞', 'done');
"""

    files["db.php"] = """<?php
// 資料庫連線（教師預先設定好，學生 require 即可）
// XAMPP 預設：帳號 root、密碼空白

$host = "localhost";
$user = "root";
$pass = "";
$dbname = "campus_system";

$conn = new mysqli($host, $user, $pass, $dbname);

if ($conn->connect_error) {
    die("連線失敗：" . $conn->connect_error);
}

// 設定編碼，避免中文亂碼
$conn->set_charset("utf8mb4");
"""

    files["form.html"] = """<!DOCTYPE html>
<html lang="zh-Hant">
<head>
  <meta charset="UTF-8">
  <title>校園報修表單</title>
  <link rel="stylesheet" href="style.css">
</head>
<body>
  <main class="wrap">
    <h1>校園報修系統</h1>
    <p class="lead">請填寫以下資料，我們會盡快處理。</p>

    <!-- 第2堂：action 先用 show.php；第3堂改成 insert.php -->
    <form action="insert.php" method="post">
      <label>報修人
        <input type="text" name="name" required>
      </label>
      <label>地點
        <input type="text" name="location" required>
      </label>
      <label>問題描述
        <textarea name="description" rows="4" required></textarea>
      </label>
      <button type="submit">送出報修</button>
    </form>

    <p><a href="list.php">查看報修列表</a></p>
  </main>
</body>
</html>
"""

    files["show.php"] = """<?php
// 第2堂：只接收並顯示，不寫入資料庫
$name = $_POST["name"] ?? "";
$location = $_POST["location"] ?? "";
$description = $_POST["description"] ?? "";
?>
<!DOCTYPE html>
<html lang="zh-Hant">
<head>
  <meta charset="UTF-8">
  <title>收到報修</title>
  <link rel="stylesheet" href="style.css">
</head>
<body>
  <main class="wrap">
    <h1>收到報修！</h1>
    <p class="big">報修人：<?php echo htmlspecialchars($name); ?></p>
    <p class="big">地點：<?php echo htmlspecialchars($location); ?></p>
    <p class="big">問題：<?php echo htmlspecialchars($description); ?></p>
    <p><a href="form.html">再送一筆</a></p>
  </main>
</body>
</html>
"""

    files["insert.php"] = """<?php
// 第3堂：接收表單並寫入資料庫（prepared statement）
require "db.php";

$name = $_POST["name"] ?? "";
$location = $_POST["location"] ?? "";
$description = $_POST["description"] ?? "";

// ★加分：空欄位檢查
if ($name === "" || $location === "" || $description === "") {
    die("請填寫完整資料。<a href='form.html'>返回</a>");
}

$sql = "INSERT INTO repairs (name, location, description, status) VALUES (?, ?, ?, 'pending')";
$stmt = $conn->prepare($sql);
$stmt->bind_param("sss", $name, $location, $description);

if ($stmt->execute()) {
    // ★★加分：寫入後跳到列表
    header("Location: list.php");
    exit;
}

echo "新增失敗：" . $conn->error;
$stmt->close();
$conn->close();
"""

    files["insert_fillblank.php"] = """<?php
// 第3堂填空版：把 ____ 填上正確關鍵字
require "____";  // 提示：db.php

$name = $_POST["____"] ?? "";
$location = $_POST["____"] ?? "";
$description = $_POST["____"] ?? "";

$sql = "INSERT INTO repairs (name, location, description, status) VALUES (?, ?, ?, 'pending')";
$stmt = $conn->prepare($sql);
$stmt->bind_param("sss", $name, $location, $description);

if ($stmt->____()) {  // 提示：execute
    echo "新增成功！<a href='list.php'>查看列表</a>";
} else {
    echo "新增失敗";
}
"""

    files["list.php"] = """<?php
// 第4堂：查詢並顯示＋簡單狀態篩選
require "db.php";

$status = $_GET["status"] ?? "all";

if ($status === "pending" || $status === "done") {
    $sql = "SELECT * FROM repairs WHERE status = ? ORDER BY created_at DESC";
    $stmt = $conn->prepare($sql);
    $stmt->bind_param("s", $status);
    $stmt->execute();
    $result = $stmt->get_result();
} else {
    $result = $conn->query("SELECT * FROM repairs ORDER BY created_at DESC");
}
?>
<!DOCTYPE html>
<html lang="zh-Hant">
<head>
  <meta charset="UTF-8">
  <title>報修列表</title>
  <link rel="stylesheet" href="style.css">
</head>
<body>
  <main class="wrap wide">
    <h1>報修管理列表</h1>
    <p><a href="form.html">新增報修</a></p>

    <form method="get" class="filter">
      <label>狀態篩選
        <select name="status" onchange="this.form.submit()">
          <option value="all" <?php if ($status === "all") echo "selected"; ?>>全部</option>
          <option value="pending" <?php if ($status === "pending") echo "selected"; ?>>待處理</option>
          <option value="done" <?php if ($status === "done") echo "selected"; ?>>已完成</option>
        </select>
      </label>
    </form>

    <table>
      <thead>
        <tr>
          <th>編號</th>
          <th>報修人</th>
          <th>地點</th>
          <th>問題描述</th>
          <th>狀態</th>
          <th>時間</th>
          <th>操作</th>
        </tr>
      </thead>
      <tbody>
      <?php while ($row = $result->fetch_assoc()): ?>
        <tr>
          <td><?php echo (int)$row["id"]; ?></td>
          <td><?php echo htmlspecialchars($row["name"]); ?></td>
          <td><?php echo htmlspecialchars($row["location"]); ?></td>
          <td><?php echo htmlspecialchars($row["description"]); ?></td>
          <td><?php echo htmlspecialchars($row["status"]); ?></td>
          <td><?php echo htmlspecialchars($row["created_at"]); ?></td>
          <td>
            <?php if ($row["status"] !== "done"): ?>
              <a href="update.php?id=<?php echo (int)$row["id"]; ?>&status=done">標示完成</a>
            <?php else: ?>
              <a href="update.php?id=<?php echo (int)$row["id"]; ?>&status=pending">改回待處理</a>
            <?php endif; ?>
          </td>
        </tr>
      <?php endwhile; ?>
      </tbody>
    </table>
  </main>
</body>
</html>
"""

    files["update.php"] = """<?php
// 第5堂：更新狀態
require "db.php";

$id = isset($_GET["id"]) ? (int)$_GET["id"] : 0;
$status = $_GET["status"] ?? "";

if ($id > 0 && ($status === "pending" || $status === "done")) {
    $sql = "UPDATE repairs SET status = ? WHERE id = ?";
    $stmt = $conn->prepare($sql);
    $stmt->bind_param("si", $status, $id);
    $stmt->execute();
    $stmt->close();
}

header("Location: list.php");
exit;
"""

    files["style.css"] = """/* 現成樣式：學生只需改標題文字或主色 --brand */
:root {
  --brand: #0f6b5c;
  --ink: #1a1a1a;
  --bg: #f4f7f6;
  --card: #ffffff;
}

* { box-sizing: border-box; }
body {
  margin: 0;
  font-family: "Noto Sans TC", "Microsoft JhengHei", sans-serif;
  background: linear-gradient(160deg, #e7f2ef 0%, var(--bg) 45%, #eef1f4 100%);
  color: var(--ink);
  min-height: 100vh;
}
.wrap {
  max-width: 640px;
  margin: 2rem auto;
  background: var(--card);
  padding: 1.5rem 1.75rem;
  border-radius: 12px;
  box-shadow: 0 8px 24px rgba(15, 107, 92, 0.08);
}
.wrap.wide { max-width: 960px; }
h1 { color: var(--brand); margin-top: 0; }
.lead { color: #4a5568; }
label { display: block; margin: 0.75rem 0; font-weight: 600; }
input, textarea, select {
  display: block;
  width: 100%;
  margin-top: 0.35rem;
  padding: 0.55rem 0.7rem;
  border: 1px solid #cfd8d6;
  border-radius: 8px;
  font: inherit;
}
button {
  margin-top: 1rem;
  background: var(--brand);
  color: #fff;
  border: 0;
  padding: 0.65rem 1.2rem;
  border-radius: 8px;
  font: inherit;
  cursor: pointer;
}
button:hover { filter: brightness(1.05); }
.big { font-size: 1.35rem; }
table {
  width: 100%;
  border-collapse: collapse;
  margin-top: 1rem;
  font-size: 0.95rem;
}
th, td {
  border-bottom: 1px solid #e2e8e6;
  padding: 0.55rem 0.4rem;
  text-align: left;
  vertical-align: top;
}
th { color: var(--brand); }
a { color: var(--brand); }
.filter { margin: 1rem 0; }
"""

    files["README.md"] = """# 校園報修 Starter Kit

給中三學生的填空／範本專案。請複製到 XAMPP 的 `htdocs/campus_repair/`。

## 檔案說明

| 檔案 | 堂次 | 用途 |
|---|---|---|
| `sql/schema.sql` | 1 | 建立資料庫與範例資料 |
| `db.php` | 3+ | 連線（教師預設） |
| `form.html` | 2–5 | 報修表單 |
| `show.php` | 2 | 只顯示、不寫入 |
| `insert.php` | 3 | 寫入資料庫（完整版） |
| `insert_fillblank.php` | 3 | 填空練習版 |
| `list.php` | 4–5 | 列表＋篩選 |
| `update.php` | 5 | 更新狀態 |
| `style.css` | 5 | 現成美化（可改 `--brand`） |

## 建議學習順序

1. 匯入 `schema.sql` 或手動建表  
2. 第2堂把 `form.html` 的 action 暫時改成 `show.php`  
3. 第3堂改回 `insert.php`  
4. 第4–5堂使用 `list.php`／`update.php`

## 安全提醒

本套件僅供課堂本機練習，不要公開部署到網際網路。
"""

    written = []
    for rel, content in files.items():
        path = STARTER / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        written.append(path)
    return written


def update_readme(extra_lines):
    readme = ROOT / "README.md"
    base = """# 校園小系統實作：數據庫與 PHP 入門（中三教材）

## 設計原則

- 全班統一做「校園報修系統」
- 每堂 1 個新概念 + 1 個可見成果
- 筆記精簡；提供 Starter Kit 與填空程式
- 教案含 VASK（價值觀／態度／技能／知識）

## 產出檔案（`output/`）

1. `01_完整教案_含VASK.docx`
2. `02_五節教學進度表.docx`
3. `03_導學案_五堂合一.docx`
4. `04_*.pptx`（合訂＋分堂簡報）
5. `05_精簡教師教案_五堂.docx`
6. `06_學生半頁筆記與任務單_*.docx`
7. `07_評分規準_作弊條_錯誤對照.docx`
8. `08_PPT大綱與逐頁要點.md`／`.docx`

## Starter Kit（`starter-kit/`）

SQL、db.php、表單、insert／list／update、填空版、style.css

## 重新產生文件

```bash
pip install -r requirements.txt
python3 scripts/generate_materials.py
python3 scripts/generate_plan_pack.py
```
"""
    readme.write_text(base + "\n".join(extra_lines) + "\n", encoding="utf-8")


def main():
    paths = []
    paths.append(build_teacher_plans())
    paths.extend(build_notes_and_tasks())
    paths.append(build_rubric_and_refs())
    paths.extend(build_ppt_outline())
    paths.extend(write_starter_kit())
    update_readme([])
    for pth in paths:
        print("Wrote", pth)


if __name__ == "__main__":
    main()
