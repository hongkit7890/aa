#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate Traditional Chinese teaching materials for Form 3 PHP/MySQL course."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt as PPt

OUT = Path(__file__).resolve().parents[1] / "output"
OUT.mkdir(parents=True, exist_ok=True)

# Brand colors (avoid purple AI cliché): teal + ink
TEAL = RGBColor(0x0F, 0x6B, 0x5C)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x4A, 0x55, 0x68)
CREAM = RGBColor(0xF7, 0xF4, 0xEF)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)


def set_run_font(run, name="Microsoft JhengHei", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True)
    return p


def add_para(doc, text, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullets(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=size)
        p.paragraph_format.space_after = Pt(2)


def set_cell_text(cell, text, bold=False, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def fill_table(table, rows):
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            set_cell_text(table.rows[i].cells[j], str(val), bold=(i == 0), size=10)


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)


# ---------------------------------------------------------------------------
# Shared curriculum content
# ---------------------------------------------------------------------------

COURSE = {
    "name": "校園小系統實作：數據庫與PHP入門",
    "audience": "中三學生（基礎能力一般）",
    "hours": "5節，每節2小時（共10小時）",
    "topic": "校園報修系統（全班統一專題）",
    "env": "XAMPP（Apache + MySQL + PHP）、phpMyAdmin、VS Code、瀏覽器、投影設備",
}

VASK = {
    "Values": [
        "尊重他人私隱：報修內容不作玩笑或公開嘲笑",
        "負責任使用科技：練習系統僅供課堂學習，不上載至公開網路",
        "誠實學習：可參考範本，但須理解並能說明自己的程式",
        "關懷校園：以科技協助解決真實校園問題",
    ],
    "Attitudes": [
        "積極嘗試：先求有、再求好，不怕出錯",
        "耐心除錯：錯誤是學習機會",
        "合作互助：願意幫忙同學、也願意請教",
        "精益求精：在完成基本功能後再美化或加分",
    ],
    "Skills": [
        "使用 phpMyAdmin 建立資料庫與資料表",
        "撰寫簡單 HTML 表單與 PHP 變數／條件",
        "以 mysqli 連接 MySQL 並執行 INSERT／SELECT／UPDATE",
        "閱讀錯誤訊息並逐步排查問題",
        "整合多頁面並向同學展示系統流程",
    ],
    "Knowledge": [
        "資料庫基本概念：資料表、欄位、記錄、主鍵",
        "系統架構：前端表單 → PHP → 數據庫 → 網頁顯示",
        "PHP 基本語法：標籤、變數、echo、if、迴圈",
        "SQL 基本指令：INSERT、SELECT、UPDATE",
        "資訊安全入門：勿直接拼接使用者輸入至 SQL",
    ],
}

LESSONS = [
    {
        "no": 1,
        "title": "系統介紹與數據庫基礎",
        "focus": "資料表／欄位／記錄／主鍵",
        "outcome": "建立 campus_system 與 repairs 表，並手動新增 3 筆資料",
        "objectives": [
            "認識校園報修專題與四格系統流程",
            "理解資料表、欄位、記錄、主鍵",
            "能在 phpMyAdmin 建立資料庫與資料表",
        ],
        "vask_focus": {
            "V": "關懷校園：科技可協助報修",
            "A": "好奇探索、認真觀察完成品",
            "S": "建立資料庫與資料表、手動新增記錄",
            "K": "資料表結構與主鍵自動遞增",
        },
        "flow": [
            ("導入", 15, "試用完成品；說明5堂後能做出類似系統"),
            ("概念講解", 30, "四格流程圖；資料表概念；示範 phpMyAdmin"),
            ("學生實作", 60, "建立資料庫／資料表；新增3筆；截圖上傳"),
            ("小結預告", 15, "回顧主鍵；預告下堂 PHP 表單"),
        ],
        "tasks": [
            "建立資料庫 campus_system",
            "建立資料表 repairs（id, name, location, description, status, created_at）",
            "手動新增 3 筆報修資料並截圖",
        ],
        "notes": [
            "欄位用「中文說明＋英文欄名」對照",
            "強調 id 為主鍵、自動遞增",
            "把資料表比喻成 Excel：一列＝一筆記錄",
        ],
    },
    {
        "no": 2,
        "title": "PHP基礎與表單處理",
        "focus": "表單 + $_POST",
        "outcome": "完成可運作的報修表單與顯示頁",
        "objectives": [
            "掌握 PHP 極簡語法",
            "能用 HTML 表單收集資料",
            "能用 $_POST 接收並顯示內容",
        ],
        "vask_focus": {
            "V": "誠實填寫練習資料，不惡意提交不當內容",
            "A": "先求有再求好；遇錯不放棄",
            "S": "製作表單、讀取 $_POST、echo 顯示",
            "K": "method/action、變數、條件判斷入門",
        },
        "flow": [
            ("回顧", 10, "資料表欄位與報修流程"),
            ("概念講解", 30, "<?php ?>、變數、echo、if；表單與 $_POST"),
            ("學生實作", 65, "form.html + show.php；測試送出顯示"),
            ("小結預告", 15, "本堂未寫入資料庫；下堂 INSERT"),
        ],
        "tasks": [
            "完成 form.html（姓名、地點、問題描述）",
            "完成 show.php 接收並大字顯示資料",
            "自我測試至少送出 2 次",
        ],
        "notes": [
            "本堂不教 for；迴圈留第4堂",
            "可用 var_dump($_POST) 除錯",
            "提醒副檔名、路徑、大小寫",
        ],
    },
    {
        "no": 3,
        "title": "PHP連接數據庫與新增資料",
        "focus": "連線 + INSERT",
        "outcome": "表單送出後資料成功存入數據庫",
        "objectives": [
            "能 require 連線檔並連接 MySQL",
            "能將表單資料 INSERT 至 repairs",
            "認識 SQL 注入風險與基本防護概念",
        ],
        "vask_focus": {
            "V": "重視安全：不把密碼寫在公開地方",
            "A": "耐心對照錯誤訊息",
            "S": "使用 mysqli／prepared statement 寫入資料",
            "K": "INSERT 語法；連線參數；安全入門",
        },
        "flow": [
            ("回顧", 10, "表單欄位對應資料表欄位"),
            ("概念講解", 30, "db.php 範本；INSERT；prepared statement 一句概念"),
            ("學生實作", 65, "insert.php 寫入；在 phpMyAdmin 核對"),
            ("小結預告", 15, "資料已進庫；下堂做列表頁"),
        ],
        "tasks": [
            "使用教師提供的 db.php",
            "把 show.php 改為 insert.php 並寫入資料",
            "確認 phpMyAdmin 出現新記錄",
        ],
        "notes": [
            "連線由教師包好，降低挫折",
            "安全只講一句＋示範一次 prepared statement",
            "預留『故意寫錯再修好』示範",
        ],
    },
    {
        "no": 4,
        "title": "查詢與顯示資料",
        "focus": "SELECT + 迴圈出表格",
        "outcome": "完成可查看／篩選報修資料的管理頁",
        "objectives": [
            "能用 SELECT 查詢資料",
            "能用 PHP 迴圈輸出 HTML 表格",
            "能加入簡單狀態篩選",
        ],
        "vask_focus": {
            "V": "尊重同學報修內容，不作公開嘲弄",
            "A": "仔細核對畫面與資料庫是否一致",
            "S": "撰寫 SELECT、迴圈、簡單篩選",
            "K": "查詢結果集；WHERE；表格呈現",
        },
        "flow": [
            ("回顧", 10, "INSERT 成功條件"),
            ("概念講解", 30, "SELECT；while 迴圈出表格；status 篩選"),
            ("學生實作", 65, "list.php；全部列表＋狀態下拉篩選"),
            ("小結預告", 15, "下堂 UPDATE 與成果發表"),
        ],
        "tasks": [
            "建立 list.php 顯示全部報修",
            "加入狀態篩選（全部／待處理／已完成）",
            "核對表格欄位與資料表一致",
        ],
        "notes": [
            "用『PHP開→表格→PHP關』三段式範本",
            "本堂是最有系統感的一頁，時間可略多",
            "篩選保持簡單，不做複雜搜尋",
        ],
    },
    {
        "no": 5,
        "title": "功能整合、更新狀態與成果發表",
        "focus": "UPDATE + 整合展示",
        "outcome": "完成可展示的報修小系統並反思",
        "objectives": [
            "能用 UPDATE 更新報修狀態",
            "整合新增、查詢、更新頁面",
            "能展示系統並反思倫理與限制",
        ],
        "vask_focus": {
            "V": "資訊倫理：隱私、不上公開網、僅供學習",
            "A": "自信展示；虛心接受回饋",
            "S": "UPDATE 狀態；整合連結；簡單美化",
            "K": "完整 CRUD 入門流程；實際部署需更多安全",
        },
        "flow": [
            ("回顧", 10, "新增與列表流程"),
            ("概念講解", 25, "UPDATE；按鈕連結；現成 CSS"),
            ("學生實作", 55, "update.php；整合頁面；套用 style.css"),
            ("成果發表", 20, "分組2分鐘：送出→列表→改狀態"),
            ("反思總結", 10, "安全、隱私、可行性三句"),
        ],
        "tasks": [
            "完成狀態更新功能",
            "整合 form／insert／list／update",
            "分組展示並完成反思三題",
        ],
        "notes": [
            "CSS 給現成檔，只改標題／主色",
            "UPDATE 可用 update.php?id=&status=",
            "強調系統僅供學習",
        ],
    },
]


# ---------------------------------------------------------------------------
# 1) Full VASK lesson plan DOCX
# ---------------------------------------------------------------------------

def build_vask_lesson_plan():
    doc = Document()
    style_doc(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("電腦科教案（含 VASK）")
    set_run_font(r, size=18, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(COURSE["name"])
    set_run_font(r, size=14, bold=True)

    add_heading(doc, "一、課程基本資料", 1)
    meta = doc.add_table(rows=6, cols=2)
    meta.style = "Table Grid"
    fill_table(
        meta,
        [
            ("項目", "內容"),
            ("課程名稱", COURSE["name"]),
            ("對象", COURSE["audience"]),
            ("課時", COURSE["hours"]),
            ("專題", COURSE["topic"]),
            ("教學環境", COURSE["env"]),
        ],
    )

    add_heading(doc, "二、課程總目標", 1)
    add_bullets(
        doc,
        [
            "理解數據庫基本概念：資料表、欄位、記錄、主鍵。",
            "掌握基本 PHP：變數、表單接收、條件、迴圈、連接數據庫。",
            "完成一套可展示的校園報修小系統。",
            "體驗「前端表單 → PHP → 數據庫 → 網頁顯示」完整流程。",
            "培養除錯、合作與資訊倫理意識。",
        ],
    )

    add_heading(doc, "三、VASK 學習目標總表", 1)
    add_para(doc, "本課程按香港課程常用之 VASK（Values／Attitudes／Skills／Knowledge）設計。", size=10)
    vask_table = doc.add_table(rows=5, cols=2)
    vask_table.style = "Table Grid"
    set_cell_text(vask_table.rows[0].cells[0], "向度", bold=True)
    set_cell_text(vask_table.rows[0].cells[1], "學習重點", bold=True)
    labels = [
        ("V 價值觀 Values", VASK["Values"]),
        ("A 態度 Attitudes", VASK["Attitudes"]),
        ("S 技能 Skills", VASK["Skills"]),
        ("K 知識 Knowledge", VASK["Knowledge"]),
    ]
    for i, (label, items) in enumerate(labels, start=1):
        set_cell_text(vask_table.rows[i].cells[0], label, bold=True)
        set_cell_text(vask_table.rows[i].cells[1], "\n".join(f"• {x}" for x in items))

    add_heading(doc, "四、評量方式", 1)
    add_para(doc, "形成性評量", bold=True)
    add_bullets(doc, ["每堂實作進度與截圖", "課堂除錯表現", "同儕互助與合作"])
    add_para(doc, "總結性評量", bold=True)
    add_bullets(
        doc,
        [
            "系統功能完整度（表單→寫入→列表→更新）",
            "資料庫設計正確性",
            "PHP 邏輯可運行",
            "展示說明與反思",
        ],
    )
    add_para(doc, "加分項目", bold=True)
    add_bullets(doc, ["空欄位驗證", "寫入後自動跳轉列表", "簡單統計／搜尋", "版面美化"])

    add_heading(doc, "五、評分規準（Rubric 摘要）", 1)
    rub = doc.add_table(rows=7, cols=5)
    rub.style = "Table Grid"
    fill_table(
        rub,
        [
            ("評分項目", "優異", "良好", "合格", "待改進"),
            (
                "資料庫設計",
                "欄位完整、主鍵正確、能清楚說明用途",
                "欄位正確、說明大致清楚",
                "能建立基本表並有資料",
                "資料表未完成或錯誤多",
            ),
            (
                "PHP程式邏輯",
                "流程順暢、能解釋關鍵程式",
                "功能可運行、解釋大致正確",
                "在協助下可運行",
                "多數功能無法運行",
            ),
            (
                "系統功能完整度",
                "新增／查詢／更新齊備且穩定",
                "三項皆有、偶有小問題",
                "完成新增＋查詢",
                "僅完成部分頁面",
            ),
            (
                "除錯與問題解決",
                "能獨立對照錯誤並修正",
                "多數能自行修正",
                "需提示才能修正",
                "遇錯即停、無排查",
            ),
            (
                "合作與展示",
                "展示清晰、分工明確、回應提問",
                "展示清楚、能說明流程",
                "能完成基本展示",
                "展示不清或未參與",
            ),
            (
                "資訊倫理與反思",
                "能連結隱私／安全／責任",
                "能說出主要倫理重點",
                "完成反思題",
                "反思空泛或未交",
            ),
        ],
    )

    add_heading(doc, "六、五堂課詳細教學設計", 1)
    for lesson in LESSONS:
        add_heading(doc, f"第{lesson['no']}堂：{lesson['title']}（120分鐘）", 2)
        add_para(doc, f"本堂焦點：{lesson['focus']}", bold=True)
        add_para(doc, f"預期成果：{lesson['outcome']}")

        add_para(doc, "教學目標", bold=True)
        add_bullets(doc, lesson["objectives"])

        add_para(doc, "本堂 VASK 焦點", bold=True)
        vt = doc.add_table(rows=5, cols=2)
        vt.style = "Table Grid"
        set_cell_text(vt.rows[0].cells[0], "向度", bold=True)
        set_cell_text(vt.rows[0].cells[1], "焦點", bold=True)
        for i, key in enumerate(["V", "A", "S", "K"], start=1):
            set_cell_text(vt.rows[i].cells[0], key, bold=True)
            set_cell_text(vt.rows[i].cells[1], lesson["vask_focus"][key])

        add_para(doc, "教學流程", bold=True)
        ft = doc.add_table(rows=len(lesson["flow"]) + 1, cols=3)
        ft.style = "Table Grid"
        set_cell_text(ft.rows[0].cells[0], "階段", bold=True)
        set_cell_text(ft.rows[0].cells[1], "時間", bold=True)
        set_cell_text(ft.rows[0].cells[2], "內容", bold=True)
        for i, (stage, mins, content) in enumerate(lesson["flow"], start=1):
            set_cell_text(ft.rows[i].cells[0], stage)
            set_cell_text(ft.rows[i].cells[1], f"{mins}分鐘")
            set_cell_text(ft.rows[i].cells[2], content)

        add_para(doc, "學生任務", bold=True)
        add_bullets(doc, lesson["tasks"])
        add_para(doc, "教師提示", bold=True)
        add_bullets(doc, lesson["notes"])

    add_heading(doc, "七、差異化與分層任務", 1)
    add_bullets(
        doc,
        [
            "★ 必做：完成本堂最低可運行成果",
            "★★ 加做：例如寫入後跳轉 list.php、空欄位提示",
            "★★★ 挑戰：簡單統計、關鍵字搜尋、借用／問卷延伸構想",
            "全班統一做報修系統，避免三專題並行造成進度分裂",
        ],
    )

    add_heading(doc, "八、附件建議", 1)
    add_bullets(
        doc,
        [
            "系統四格流程圖",
            "資料表欄位設計對照表",
            "Starter Kit：db.php、form.html、insert.php、list.php、update.php、style.css",
            "學生導學案與任務單",
            "錯誤對照表與一頁作弊條",
            "評分規準表",
        ],
    )

    add_heading(doc, "九、安全與倫理提醒（教師）", 1)
    add_bullets(
        doc,
        [
            "練習系統僅在本機／內網使用",
            "勿收集身分證等敏感個資",
            "帳號密碼勿投影過久；示範後提醒更改",
            "實際公開部署需額外安全措施（本課程不涵蓋）",
        ],
    )

    path = OUT / "01_完整教案_含VASK.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 2) Progress schedule DOCX
# ---------------------------------------------------------------------------

def build_progress_schedule():
    doc = Document()
    style_doc(doc)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("教學進度表")
    set_run_font(r, size=18, bold=True)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(f"{COURSE['name']}｜{COURSE['hours']}")
    set_run_font(r, size=12, bold=True)

    add_para(doc, f"對象：{COURSE['audience']}　　專題：{COURSE['topic']}")

    add_heading(doc, "一、總覽進度表", 1)
    table = doc.add_table(rows=6, cols=6)
    table.style = "Table Grid"
    headers = ["堂次", "主題", "本堂焦點", "預期產出", "VASK 主軸", "評核證據"]
    for j, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], h, bold=True)
    overview = [
        (
            "1",
            "系統介紹與數據庫基礎",
            "資料表概念",
            "campus_system + repairs + 3筆資料截圖",
            "K為主，兼V關懷校園",
            "資料庫截圖",
        ),
        (
            "2",
            "PHP基礎與表單處理",
            "表單與$_POST",
            "form.html + show.php 可運行",
            "S表單處理，A先求有",
            "送出顯示截圖",
        ),
        (
            "3",
            "連接數據庫與新增",
            "連線與INSERT",
            "insert.php 成功寫入",
            "S寫入，V安全意識",
            "phpMyAdmin新記錄",
        ),
        (
            "4",
            "查詢與顯示資料",
            "SELECT與列表",
            "list.php + 狀態篩選",
            "S查詢顯示，V尊重內容",
            "管理頁截圖",
        ),
        (
            "5",
            "整合、更新與發表",
            "UPDATE與展示",
            "可展示完整小系統",
            "全程VASK整合",
            "展示＋反思單",
        ),
    ]
    for i, row in enumerate(overview, start=1):
        for j, val in enumerate(row):
            set_cell_text(table.rows[i].cells[j], val)

    add_heading(doc, "二、每堂分鐘進度（120分鐘）", 1)
    for lesson in LESSONS:
        add_para(doc, f"第{lesson['no']}堂：{lesson['title']}", bold=True)
        ft = doc.add_table(rows=len(lesson["flow"]) + 1, cols=4)
        ft.style = "Table Grid"
        for j, h in enumerate(["順序", "階段", "時間", "教學重點"]):
            set_cell_text(ft.rows[0].cells[j], h, bold=True)
        for i, (stage, mins, content) in enumerate(lesson["flow"], start=1):
            set_cell_text(ft.rows[i].cells[0], str(i))
            set_cell_text(ft.rows[i].cells[1], stage)
            set_cell_text(ft.rows[i].cells[2], f"{mins}′")
            set_cell_text(ft.rows[i].cells[3], content)
        doc.add_paragraph()

    add_heading(doc, "三、里程碑檢查點", 1)
    mt = doc.add_table(rows=6, cols=3)
    mt.style = "Table Grid"
    fill_table(
        mt,
        [
            ("檢查點", "時間", "通過標準"),
            ("CP1", "第1堂結束", "資料表存在且有≥3筆記錄"),
            ("CP2", "第2堂結束", "表單送出後網頁顯示正確欄位"),
            ("CP3", "第3堂結束", "新資料出現在資料庫"),
            ("CP4", "第4堂結束", "列表頁可顯示並可篩選"),
            ("CP5", "第5堂結束", "可更新狀態並完成展示／反思"),
        ],
    )

    add_heading(doc, "四、備課與物資時程", 1)
    add_bullets(
        doc,
        [
            "開課前：安裝測試 XAMPP；準備完成品示範；複製 Starter Kit",
            "每堂前：檢查投影、示例資料庫、本堂填空程式紙",
            "第5堂前：安排展示順序與計時；印反思單",
        ],
    )

    path = OUT / "02_五節教學進度表.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 3) Study guide (導學案) DOCX
# ---------------------------------------------------------------------------

def build_study_guide():
    doc = Document()
    style_doc(doc)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("學生導學案")
    set_run_font(r, size=18, bold=True)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(f"{COURSE['name']}（五堂合一）")
    set_run_font(r, size=13, bold=True)

    add_para(doc, "姓名：__________　　班別：____　　學號：________　　組別：____")
    add_para(
        doc,
        "使用說明：每堂先看「我要學會」，再跟步驟做；完成後勾選檢查清單。筆記已精簡，以動手成功為先。",
        size=10,
    )

    add_heading(doc, "課程地圖（請記住這4步）", 1)
    add_para(doc, "① 填表單  →  ② PHP 處理  →  ③ 存進資料庫  →  ④ 網頁列表顯示", bold=True)

    add_heading(doc, "資料表速查（全程使用）", 1)
    dt = doc.add_table(rows=7, cols=3)
    dt.style = "Table Grid"
    fill_table(
        dt,
        [
            ("欄位名", "中文意思", "例子"),
            ("id", "編號（主鍵，自動加1）", "1"),
            ("name", "報修人", "陳同學"),
            ("location", "地點", "3樓電腦室"),
            ("description", "問題描述", "投影機無法開機"),
            ("status", "狀態", "pending／done"),
            ("created_at", "建立時間", "自動或手動填"),
        ],
    )

    # Per-lesson pages
    guide_extras = {
        1: {
            "warmup": "試用老師的完成品：送出一筆報修，看看列表有沒有出現。",
            "mini_notes": [
                "資料庫像一本練習簿；資料表像其中一頁表格。",
                "欄位＝欄標題；記錄＝一橫列資料；主鍵＝永不重複的編號。",
                "本課資料庫名：campus_system；資料表名：repairs。",
            ],
            "steps": [
                "打開 XAMPP，啟動 Apache 與 MySQL。",
                "打開 phpMyAdmin，建立資料庫 campus_system。",
                "在 campus_system 建立資料表 repairs，加入指定欄位。",
                "把 id 設為主鍵（PRIMARY）並勾選 A_I（自動遞增）。",
                "用「插入」手動新增 3 筆報修資料。",
                "截圖：資料表結構＋3筆資料，上傳課堂平台。",
            ],
            "checklist": [
                "我能說出資料表／欄位／記錄／主鍵的意思",
                "campus_system 與 repairs 已建立",
                "至少有 3 筆資料",
                "已截圖上傳",
            ],
            "reflect": "用自己的話：主鍵為什麼重要？________________",
        },
        2: {
            "warmup": "回想：報修表單需要哪些欄位？對應資料表哪幾欄？",
            "mini_notes": [
                "PHP 程式寫在 <?php ... ?> 之間。",
                "變數以 $ 開頭，例如 $name。",
                "表單 method=\"post\"，action 指向處理頁（如 show.php）。",
                "用 $_POST['欄位name'] 接收資料，再用 echo 顯示。",
            ],
            "steps": [
                "在 htdocs 建立資料夾 campus_repair。",
                "建立 form.html：姓名、地點、問題描述、送出按鈕。",
                "建立 show.php：接收 $_POST 並大字顯示三個欄位。",
                "用瀏覽器打開 form.html，送出測試。",
                "若空白：檢查 input 的 name、檔名、路徑。",
            ],
            "checklist": [
                "表單可以打開",
                "送出後 show.php 顯示正確內容",
                "我知道 $_POST 的用途",
            ],
            "reflect": "本堂資料有沒有進資料庫？為什麼？________________",
        },
        3: {
            "warmup": "把 show.php 顯示的三個值，對應到 repairs 的欄位名。",
            "mini_notes": [
                "使用老師提供的 db.php（裡面已寫好連線）。",
                "insert.php：接收表單 → 連接資料庫 → INSERT。",
                "安全一句：不要把使用者輸入直接拼進 SQL；用 prepared statement。",
                "成功後到 phpMyAdmin 刷新，應看到新列。",
            ],
            "steps": [
                "複製 db.php 到專案資料夾。",
                "把表單 action 改為 insert.php。",
                "完成 insert.php（可填空）：連線、預備 SQL、綁定參數、執行。",
                "送出表單，到 phpMyAdmin 確認新資料。",
                "（加分）成功後自動跳到 list.php。",
            ],
            "checklist": [
                "表單改指向 insert.php",
                "送出後資料庫多一筆",
                "我能指出 INSERT 大概在做什麼",
            ],
            "reflect": "若連線失敗，你會先檢查哪兩件事？________________",
        },
        4: {
            "warmup": "在 phpMyAdmin 數一數目前有幾筆報修。",
            "mini_notes": [
                "SELECT * FROM repairs; 可取出全部資料。",
                "PHP 用迴圈（while）一列一列輸出成 HTML 表格。",
                "篩選：WHERE status = 'pending' 或 'done'。",
                "排版口訣：PHP 開 → 表格 → PHP 關。",
            ],
            "steps": [
                "建立 list.php，連接資料庫。",
                "查詢全部報修並以表格顯示。",
                "加入狀態下拉：全部／待處理／已完成。",
                "測試篩選是否正確。",
                "（加分）依時間新到舊排序。",
            ],
            "checklist": [
                "列表能顯示資料庫中的資料",
                "篩選可以使用",
                "表格欄位齊全易讀",
            ],
            "reflect": "列表頁為什麼讓系統『看起來像真的系統』？________________",
        },
        5: {
            "warmup": "從 form 走到 list，畫出你的檔案連結關係。",
            "mini_notes": [
                "UPDATE repairs SET status='done' WHERE id=?;",
                "可用連結：update.php?id=1&status=done",
                "套用老師的 style.css，只改標題與主色即可。",
                "倫理三句：少收集敏感資料；不上公開網；尊重他人內容。",
            ],
            "steps": [
                "完成 update.php，可把 pending 改成 done（或改回）。",
                "在 list.php 每一列加「標示完成」連結。",
                "整合導覽：表單／列表互相可點。",
                "套用 CSS，檢查手機／電腦瀏覽器基本可讀。",
                "分組展示（2分鐘）＋完成反思三題。",
            ],
            "checklist": [
                "可以更新狀態",
                "頁面能互相跳轉",
                "完成展示與反思",
            ],
            "reflect": "若真的給全校用，還缺哪些安全措施？________________",
        },
    }

    for lesson in LESSONS:
        extra = guide_extras[lesson["no"]]
        doc.add_page_break()
        add_heading(doc, f"第{lesson['no']}堂　{lesson['title']}", 1)
        add_para(doc, f"本堂焦點：{lesson['focus']}", bold=True)
        add_para(doc, f"下課前成功長這樣：{lesson['outcome']}", bold=True)

        add_para(doc, "【我要學會】", bold=True)
        add_bullets(doc, lesson["objectives"])

        add_para(doc, "【VASK 小提醒】", bold=True)
        add_bullets(
            doc,
            [
                f"V：{lesson['vask_focus']['V']}",
                f"A：{lesson['vask_focus']['A']}",
                f"S：{lesson['vask_focus']['S']}",
                f"K：{lesson['vask_focus']['K']}",
            ],
        )

        add_para(doc, "【課前熱身】", bold=True)
        add_para(doc, extra["warmup"])

        add_para(doc, "【精簡筆記】", bold=True)
        add_bullets(doc, extra["mini_notes"])

        add_para(doc, "【實作步驟】", bold=True)
        for i, step in enumerate(extra["steps"], 1):
            add_para(doc, f"{i}. {step}")

        add_para(doc, "【自我檢查】", bold=True)
        for c in extra["checklist"]:
            add_para(doc, f"☐ {c}")

        add_para(doc, "【一分鐘反思】", bold=True)
        add_para(doc, extra["reflect"])

        add_para(doc, "【教師／自己評】　★完成　／　需協助：__________", size=10)

    doc.add_page_break()
    add_heading(doc, "附錄A：常見錯誤對照（精簡）", 1)
    et = doc.add_table(rows=7, cols=2)
    et.style = "Table Grid"
    fill_table(
        et,
        [
            ("現象", "先檢查"),
            ("網頁打不開", "Apache 是否啟動；網址路徑是否正確"),
            ("送出後空白／Notice", "input 的 name 是否與 $_POST 鍵名一致"),
            ("連線失敗", "MySQL 是否啟動；db.php 帳密資料庫名"),
            ("INSERT 後看不到", "是否選錯資料庫；是否真的執行成功"),
            ("列表空白", "資料表是否有資料；SQL 是否寫錯表名"),
            ("中文亂碼", "檔案與資料庫編碼是否為 utf8mb4／UTF-8"),
        ],
    )

    add_heading(doc, "附錄B：第5堂反思三題", 1)
    add_para(doc, "1. 我學會的一個資料庫概念：________________")
    add_para(doc, "2. 我遇過的一個錯誤，以及如何解決：________________")
    add_para(doc, "3. 關於私隱／安全，我記住的一句：________________")

    add_heading(doc, "附錄C：一頁作弊條", 1)
    add_para(doc, "PHP：<?php  $x=1;  echo $x;  if(...){ }  ?>", bold=True)
    add_para(doc, "接收表單：$name = $_POST['name'];")
    add_para(doc, "SQL：INSERT INTO repairs (name,location,description,status) VALUES (?,?,?,?)")
    add_para(doc, "SQL：SELECT * FROM repairs;")
    add_para(doc, "SQL：UPDATE repairs SET status='done' WHERE id=?;")

    path = OUT / "03_導學案_五堂合一.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 4) PPTX (five lessons in one deck, ~8-12 slides each)
# ---------------------------------------------------------------------------

def _set_run(run, size=20, bold=False, color=INK, font="Microsoft JhengHei"):
    run.font.size = PPt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=CREAM):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_bar(slide, text):
    shape = slide.shapes.add_shape(
        1, Inches(0), Inches(0), Inches(13.333), Inches(1.0)
    )  # rectangle
    shape.fill.solid()
    shape.fill.fore_color.rgb = TEAL
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "  " + text
    _set_run(run, size=26, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))


def add_body_box(slide, lines, left=0.6, top=1.3, width=12.0, height=5.8, size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = 0
        run = p.add_run()
        run.text = line
        _set_run(run, size=size, bold=False, color=INK)
        p.space_after = PPt(8)
    return box


def new_slide(prs, title, lines, footer=None, size=20):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_bg(slide)
    add_title_bar(slide, title)
    add_body_box(slide, lines, size=size)
    if footer:
        box = slide.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(12), Inches(0.4))
        tf = box.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = footer
        _set_run(run, size=12, color=MUTED)
    return slide


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ---- Cover ----
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = TEAL
    shape.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(3.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = COURSE["name"]
    _set_run(run, size=36, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    run2.text = "中三｜5節×2小時｜校園報修系統"
    _set_run(run2, size=22, color=RGBColor(0xD7, 0xF0, 0xEA))
    p3 = tf.add_paragraph()
    run3 = p3.add_run()
    run3.text = "數據庫與 PHP 入門｜含 VASK 學習目標"
    _set_run(run3, size=18, color=RGBColor(0xD7, 0xF0, 0xEA))

    new_slide(
        prs,
        "課程目標",
        [
            "1. 理解資料表、欄位、記錄、主鍵",
            "2. 掌握基本 PHP：變數、表單、條件、迴圈、連資料庫",
            "3. 完成校園報修小系統（可展示）",
            "4. 體驗：表單 → PHP → 數據庫 → 網頁顯示",
            "5. 培養除錯、合作與資訊倫理",
        ],
        footer="教師提示：全班統一專題，降低分叉",
    )

    new_slide(
        prs,
        "VASK 總覽",
        [
            "V 價值觀：尊重私隱、負責任使用科技、誠實學習、關懷校園",
            "A 態度：積極嘗試、耐心除錯、合作互助、先求有再求好",
            "S 技能：建表、表單、INSERT／SELECT／UPDATE、除錯、展示",
            "K 知識：資料庫概念、系統四步流程、PHP／SQL 基礎、安全入門",
        ],
        footer="每堂都會標出本堂 VASK 焦點",
        size=18,
    )

    new_slide(
        prs,
        "記住系統四步",
        [
            "① 使用者填前端表單",
            "② PHP 接收並處理資料",
            "③ 資料存入 MySQL 資料庫",
            "④ 再用網頁把資料顯示出來",
            "",
            "這四步＝你們五堂課要拼起來的完整流程",
        ],
        footer="建議視覺：四格流程圖",
    )

    # Lesson 1 slides
    L1 = [
        (
            "第1堂｜系統介紹與數據庫基礎",
            [
                "本堂焦點：資料表、欄位、記錄、主鍵",
                "下課成功：campus_system + repairs + 3筆資料",
                "時間：120分鐘",
            ],
        ),
        (
            "第1堂｜學習目標與 VASK",
            [
                "目標：認識專題；理解資料表概念；能建立資料庫／表",
                "V：關懷校園——科技可協助報修",
                "A：好奇探索、認真觀察完成品",
                "S：建立資料庫與資料表、手動新增",
                "K：資料表結構與主鍵自動遞增",
            ],
        ),
        (
            "第1堂｜先玩完成品（導入）",
            [
                "請同學試用報修系統：送出一筆",
                "問：資料去了哪裡？如何再被看見？",
                "宣布：五堂後你也能做出類似系統",
            ],
            "教師：控制在15分鐘，留下成就感",
        ),
        (
            "第1堂｜四格流程圖",
            [
                "填表單 → PHP → 資料庫 → 列表顯示",
                "今天先攻『資料庫』這一格",
                "先把資料放進『倉庫』，後面再學搬運",
            ],
        ),
        (
            "第1堂｜資料表像 Excel",
            [
                "資料表 ≈ 一張表",
                "欄位 ≈ 欄標題（報修人、地點…）",
                "記錄 ≈ 一橫列資料",
                "主鍵 id ≈ 永不重複的編號（自動加1）",
            ],
        ),
        (
            "第1堂｜欄位設計表",
            [
                "id｜編號｜主鍵 A_I",
                "name｜報修人",
                "location｜地點",
                "description｜問題描述",
                "status｜狀態（pending／done）",
                "created_at｜建立時間",
            ],
            "建議視覺：對照表投影",
        ),
        (
            "第1堂｜學生實作",
            [
                "1. 啟動 XAMPP（Apache + MySQL）",
                "2. phpMyAdmin 建立 campus_system",
                "3. 建立 repairs 並設定主鍵",
                "4. 手動新增 3 筆資料",
                "5. 截圖上傳",
            ],
        ),
        (
            "第1堂｜檢查與預告",
            [
                "☐ 資料庫與資料表完成",
                "☐ 至少 3 筆資料",
                "☐ 能用自己的話說主鍵",
                "下堂：PHP 表單——先顯示，不寫入",
            ],
        ),
    ]

    L2 = [
        (
            "第2堂｜PHP基礎與表單處理",
            [
                "本堂焦點：HTML 表單 + $_POST",
                "下課成功：form.html + show.php 可運行",
                "本堂不連資料庫、不教 for 迴圈",
            ],
        ),
        (
            "第2堂｜學習目標與 VASK",
            [
                "目標：極簡 PHP；表單收集；$_POST 顯示",
                "V：誠實填寫練習資料",
                "A：先求有再求好",
                "S：做表單、讀 $_POST、echo",
                "K：method／action、變數、if 入門",
            ],
        ),
        (
            "第2堂｜PHP 只要這幾句",
            [
                "<?php  ...  ?>",
                "$name = \"小明\";",
                "echo $name;",
                "if ($name != \"\") { echo \"有資料\"; }",
            ],
            "教師：語法極簡，現場敲一遍",
        ),
        (
            "第2堂｜表單三要素",
            [
                "method=\"post\"：用 POST 送資料",
                "action=\"show.php\"：送給誰處理",
                "input 的 name：PHP 用這個名字取值",
            ],
        ),
        (
            "第2堂｜$_POST 怎麼用",
            [
                "$name = $_POST['name'];",
                "$location = $_POST['location'];",
                "$description = $_POST['description'];",
                "echo \"收到：$name\";",
            ],
        ),
        (
            "第2堂｜學生實作",
            [
                "建立 form.html（姓名、地點、描述）",
                "建立 show.php 大字顯示三欄",
                "測試送出至少 2 次",
                "卡住就 var_dump($_POST)",
            ],
        ),
        (
            "第2堂｜常見錯誤三招",
            [
                "副檔名是否 .php／.html",
                "路徑／資料夾是否正確",
                "name 與 $_POST['...'] 是否一致（大小寫）",
            ],
        ),
        (
            "第2堂｜檢查與預告",
            [
                "☐ 表單可送出",
                "☐ 顯示頁內容正確",
                "問：資料有進資料庫嗎？還沒！",
                "下堂：INSERT 真正存進去",
            ],
        ),
    ]

    L3 = [
        (
            "第3堂｜連接數據庫與新增資料",
            [
                "本堂焦點：連線 + INSERT",
                "下課成功：表單資料出現在 phpMyAdmin",
                "使用教師提供的 db.php",
            ],
        ),
        (
            "第3堂｜學習目標與 VASK",
            [
                "目標：連 MySQL；INSERT；認識基本防護",
                "V：重視帳密與安全",
                "A：耐心對照錯誤訊息",
                "S：mysqli／prepared statement 寫入",
                "K：INSERT；連線參數；安全入門",
            ],
        ),
        (
            "第3堂｜資料流對照",
            [
                "表單 name  → 資料表 name",
                "表單 location → location",
                "表單 description → description",
                "status 先固定為 pending",
            ],
        ),
        (
            "第3堂｜db.php 做什麼",
            [
                "負責：主機、帳號、密碼、資料庫名",
                "學生：require 'db.php'; 即可",
                "好處：連線問題集中處理，少挫折",
            ],
            "教師：不要投影密碼過久",
        ),
        (
            "第3堂｜INSERT（概念）",
            [
                "INSERT INTO repairs (name, location, description, status)",
                "VALUES (?, ?, ?, 'pending');",
                "？＝稍後填入的使用者輸入",
            ],
        ),
        (
            "第3堂｜安全一句話",
            [
                "不要把使用者輸入直接拼進 SQL",
                "使用 prepared statement（預處理）",
                "本課：示範一次＋填空完成即可",
            ],
        ),
        (
            "第3堂｜學生實作",
            [
                "表單 action 改 insert.php",
                "完成寫入程式（填空）",
                "送出後到 phpMyAdmin 刷新核對",
                "加分：成功後跳轉 list.php",
            ],
        ),
        (
            "第3堂｜除錯示範＋預告",
            [
                "教師故意寫錯 → 讀錯誤 → 修好",
                "☐ 新記錄成功進庫",
                "下堂：把資料用表格列出來（最有系統感）",
            ],
        ),
    ]

    L4 = [
        (
            "第4堂｜查詢與顯示資料",
            [
                "本堂焦點：SELECT + 迴圈出表格",
                "下課成功：list.php 管理頁＋狀態篩選",
                "這頁會讓作品『看起來像真系統』",
            ],
        ),
        (
            "第4堂｜學習目標與 VASK",
            [
                "目標：SELECT；迴圈輸出；簡單篩選",
                "V：尊重同學報修內容",
                "A：核對畫面與資料庫一致",
                "S：查詢、迴圈、篩選",
                "K：結果集、WHERE、表格呈現",
            ],
        ),
        (
            "第4堂｜SELECT",
            [
                "全部：SELECT * FROM repairs;",
                "待處理：... WHERE status='pending';",
                "已完成：... WHERE status='done';",
            ],
        ),
        (
            "第4堂｜PHP 迴圈出表格",
            [
                "口訣：PHP 開 → 表格 → PHP 關",
                "while 一列資料 → 輸出一個 <tr>",
                "欄位對應 <td>",
            ],
            "建議視覺：程式＋結果對照截圖",
        ),
        (
            "第4堂｜篩選 UI（保持簡單）",
            [
                "下拉選單：全部／待處理／已完成",
                "送出後帶 status 參數再查詢",
                "不做複雜關鍵字搜尋（可作加分）",
            ],
        ),
        (
            "第4堂｜學生實作",
            [
                "建立 list.php",
                "顯示全部報修為表格",
                "加入狀態篩選並測試",
                "加分：依時間排序",
            ],
        ),
        (
            "第4堂｜檢查與預告",
            [
                "☐ 列表有資料",
                "☐ 篩選有效",
                "下堂：UPDATE 改狀態＋整合展示＋倫理反思",
            ],
        ),
    ]

    L5 = [
        (
            "第5堂｜整合、更新與成果發表",
            [
                "本堂焦點：UPDATE + 整合 + 展示",
                "下課成功：可展示的報修小系統",
                "CSS 用現成檔，專心功能與表達",
            ],
        ),
        (
            "第5堂｜學習目標與 VASK",
            [
                "目標：UPDATE；整合頁面；展示與反思",
                "V：隱私、不上公開網、僅供學習",
                "A：自信展示、虛心回饋",
                "S：更新狀態、整合連結、簡單美化",
                "K：完整流程；部署需更多安全",
            ],
        ),
        (
            "第5堂｜UPDATE 狀態",
            [
                "UPDATE repairs SET status='done' WHERE id=?;",
                "列表每列加連結：update.php?id=...&status=done",
                "更新後回到 list.php",
            ],
        ),
        (
            "第5堂｜整合檢查清單",
            [
                "form → insert → 資料庫",
                "list 顯示全部／篩選",
                "update 可改狀態",
                "頁面之間有導覽連結",
            ],
        ),
        (
            "第5堂｜美化原則",
            [
                "套用 style.css",
                "只改：標題文字、主色",
                "可讀、清楚比花俏重要",
            ],
        ),
        (
            "第5堂｜展示方式（2分鐘）",
            [
                "1. 示範送出一筆報修",
                "2. 在列表指出新資料",
                "3. 把它標示為已完成",
                "加一句：資料庫如何設計／哪個 PHP 最關鍵",
            ],
        ),
        (
            "第5堂｜資訊倫理三句",
            [
                "1. 少收集敏感個人資料",
                "2. 練習系統不要公開上網",
                "3. 尊重他人的報修內容",
            ],
        ),
        (
            "第5堂｜課程總結",
            [
                "你們完成了：建表 → 表單 → 寫入 → 列表 → 更新",
                "這就是小型資訊系統的基本循環",
                "錯誤不可怕；看懂錯誤才成長",
                "謝謝努力——去展示你們的系統！",
            ],
        ),
    ]

    lesson_packs = [
        (1, "系統介紹與數據庫基礎", L1),
        (2, "PHP基礎與表單處理", L2),
        (3, "PHP連接數據庫與新增資料", L3),
        (4, "查詢與顯示資料", L4),
        (5, "功能整合更新狀態與成果發表", L5),
    ]

    def add_lesson_block(target_prs, blocks):
        for item in blocks:
            if len(item) == 2:
                title, lines = item
                footer = None
            else:
                title, lines, footer = item
            new_slide(target_prs, title, lines, footer=footer, size=20)

    def add_section_divider(target_prs, title):
        slide = target_prs.slides.add_slide(target_prs.slide_layouts[6])
        add_bg(slide, TEAL)
        box = slide.shapes.add_textbox(Inches(1), Inches(2.8), Inches(11), Inches(2))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title
        _set_run(run, size=32, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    # Combined deck
    for num, title, blocks in lesson_packs:
        add_section_divider(prs, f"第{num}堂：{title}")
        add_lesson_block(prs, blocks)

    new_slide(
        prs,
        "附件與延伸",
        [
            "附件：欄位表、Starter Kit、錯誤對照、評分規準",
            "加分：空欄位檢查、跳轉列表、統計、美化",
            "延伸專題構想：物品借用／校園問卷（學有餘力）",
        ],
        footer="教材產出：教案｜進度表｜導學案｜本簡報",
    )

    combined = OUT / "04_五堂教學簡報_合訂.pptx"
    prs.save(combined)

    # Per-lesson decks
    per_lesson_paths = []
    for num, title, blocks in lesson_packs:
        one = Presentation()
        one.slide_width = Inches(13.333)
        one.slide_height = Inches(7.5)
        add_section_divider(one, f"第{num}堂：{title}")
        add_lesson_block(one, blocks)
        path = OUT / f"04_第{num}堂_{title}.pptx"
        one.save(path)
        per_lesson_paths.append(path)

    return [combined] + per_lesson_paths


def main():
    paths = [
        build_vask_lesson_plan(),
        build_progress_schedule(),
        build_study_guide(),
    ]
    paths.extend(build_pptx())
    readme = OUT.parent / "README.md"
    readme.write_text(
        """# 校園小系統實作：數據庫與 PHP 入門（中三教材）

## 產出檔案（`output/`）

1. `01_完整教案_含VASK.docx` — 含 Values／Attitudes／Skills／Knowledge 的完整教案
2. `02_五節教學進度表.docx` — 五節總覽、分鐘進度、里程碑
3. `03_導學案_五堂合一.docx` — 學生導學案（精簡筆記＋步驟＋檢查清單）
4. `04_五堂教學簡報_合訂.pptx` — 五堂合訂簡報
5. `04_第1堂_…` 至 `04_第5堂_….pptx` — 分堂簡報（每堂約 8–10 頁）

## 重新產生

```bash
python3 scripts/generate_materials.py
```

## 課程設計原則

- 全班統一做「校園報修系統」
- 每堂 1 個新概念 + 1 個可見成果
- 筆記精簡、Starter Kit 填空、先求有再求好
- 教案含完整 VASK（價值觀／態度／技能／知識）
""",
        encoding="utf-8",
    )
    for p in paths:
        print("Wrote", p)


if __name__ == "__main__":
    main()
